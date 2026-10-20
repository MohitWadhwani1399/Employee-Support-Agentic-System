from pathlib import Path
from langchain_community.document_loaders.base_o365 import CHUNK_SIZE
from langchain_core.documents import Document
from typing import Iterable
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from docx import Document as DocxDocument

SUPPORTED_DOCS = {'.pdf','.txt','.md','.docx'}

def load_file(path:Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == '.pdf':
        return PyPDFLoader(str(path)).load()
    if suffix in {'.txt','.md'}:
        return TextLoader(str(path),encoding='utf-8').load()
    if suffix== '.docx':
        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return [Document(page_content=text, metadata={"source": str(path)})]
    raise ValueError(f"Unsupported file type: {suffix}")

def chunk_documents(docs:Iterable[Document]) -> list[Document]:
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=200,
        chunk_overlap=50,
        length_function=len,
        add_start_index=True
    )
    return text_splitter.split_documents(list(docs))

